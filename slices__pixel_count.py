import numpy as np
import functions as func
import time
import openpyxl

directories = ['1_cuboid/cube_rand_axis/', '1_cuboid/cube_z/', '2_filled_cuboid/f_cub_x/', '2_filled_cuboid/f_cub_z/', '3_cylinder/cyl_rand_axis/', '3_cylinder/cyl_x/', '3_cylinder/cyl_z/', '4_filled_cylinder/f_cyl_axis_unal/', '4_filled_cylinder/f_cyl_x/', '4_filled_cylinder/f_cyl_z/', '5_torus/torus_rand_axis/', '5_torus/torus_x/', '5_torus/torus_z/', '6_clot_1/', '7_clot_2/', '7_clot_2_pt1/', '7_clot_2_pt2/', '7_clot_2_pt3/', '8_clot_glob/', '9_mini_clot/']
nbr_images = [141, 200, 195, 198, 157, 161, 200, 183, 197, 200, 95, 149, 39, 33, 33, 33, 33, 33, 33, 33]
masks_paths = ['slices/6_clot_1/masks/', 'slices/7_clot_2/masks/', 'slices/7_clot_2_pt1/masks/', 'slices/7_clot_2_pt2/masks/', 'slices/7_clot_2_pt3/masks/', 'slices/8_clot_glob/masks/', 'slices/9_mini_clot/masks/']
masks_names = ['Crop_1', 'Crop_2', 'Crop_2', 'Crop_2', 'Crop_2', '2019_5 MSB', '0_miniclot']
ref_colors = [['#e65a49', '#ffffff'], ['#e65a49', '#ffffff'], ['#e3bb37', '#e65a49',  '#f08a8c', '#ffffff'], ['#e3bb37', '#e65a49',  '#f08a8c', '#ffffff'], ['#e65a49', '#ffffff'], ['#e65a49', '#ffffff'], ['#e65a49', '#ffffff'], ['#e3bb37', '#e65a49', '#c78637', '#f08a8c', '#ffffff'], ['#e3bb37', '#e65a49', '#c78637', '#f08a8c', '#e88d76', '#ffffff'], ['#e3bb37', '#e65a49', '#c78637', '#f08a8c', '#e88d76', '#ffffff'], ['#e65a49', '#ffffff'], ['#e65a49', '#ffffff'],  ['#e65a49', '#ffffff'], ['#d76c3a', '#f58b83', '#f6b354', '#fabeb3', '#ffffff'], ['#d76c3a', '#f58b83', '#f6b354', '#fabeb3', '#ffffff'], ['#d76c3a', '#f58b83', '#f6b354', '#fabeb3', '#ffffff'], ['#d76c3a', '#f58b83', '#f6b354', '#fabeb3', '#ffffff'], ['#d76c3a', '#f58b83', '#f6b354', '#fabeb3', '#ffffff'], ['#d76c3a', '#f58b83', '#f6b354', '#fabeb3', '#ffffff'], ['#d76c3a', '#f58b83', '#f6b354', '#fabeb3', '#ffffff']]
nbr_cluster = [2, 2, 4, 4, 2, 2, 2, 5, 6, 6, 2, 2, 2, 5, 5, 5, 5, 5, 5, 5]
i = 17 # index of the targeted dir in directories (between 0 and 19)
# 0-1 : cuboid
# 2-3 : filled cuboid
# 4-6 : cylinder
# 7-9 : filled cylinder
# 10-12 : torus
# 13-19 : clots


#color_mode = 'clust' # color clustering
color_mode = 'class' # color classification

#color_mode = 'd_clust' # double color clustering

merging_col = True

# excel file creation/loading
excel_file_exists = True

column_letters = ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P', 'Q', 'R', 'S', 'T', 'U', 'V', 'W', 'X', 'Y', 'Z']

if i <= 12 :
    idx = [j for j in range(len(directories[i])) if directories[i].startswith('/', j)]
    excel_file = 'slices/'+directories[i][0]+'_'+directories[i][idx[-2]+1:idx[-1]]+'_pixel_count'
else :
    excel_file = 'slices/'+directories[i][:-1]+'_pixel_count'


if excel_file_exists :
    excel_wb = openpyxl.load_workbook(filename=excel_file+'.xlsx')
else :
    excel_wb =openpyxl.Workbook()
    sheet = excel_wb.active
    sheet.title='Color clustering'
    excel_wb.create_sheet('Color classification')
    #excel_wb.create_sheet('Double color clustering')

# pre alignement
if color_mode == 'clust' :
    directories[i] = directories[i]+'Col_clust/'
    images = func.load_slices('slices/'+directories[i]+'2_img_col_clust', nbr_images[i])
    sheet = excel_wb['Color clustering']

elif color_mode == 'class' :
    directories[i] = directories[i]+'Col_class/'
    images = func.load_slices('slices/'+directories[i]+'2_img_col_class', nbr_images[i])
    sheet = excel_wb['Color classification']

'''
elif color_mode == 'd_clust' :
    images = func.load_slices('slices/'+directories[i]+'2_img_d_col_clust', nbr_images[i])
    directories[i] = directories[i]+'D_col_clust/'
    sheet = excel_wb['Double color clustering']
'''

counts = func.pixel_count(images)

start_r = 4
end_r = start_r + nbr_images[i] -1
start_c = 3

# filename
s = sheet.cell(row = start_r -1, column = 1)
s.value = directories[i][:-1]
sheet.merge_cells(start_row=start_r-1, start_column= 1, end_row=end_r+1, end_column=1)

# image number
im = sheet.cell(row = 1, column = 2)
im.value = 'image'
sheet.merge_cells(start_row=1, start_column= 2, end_row=3, end_column=2)
for r in range(start_r, end_r+1) :
    c = sheet.cell(row = r, column = 2)
    c. value = r-start_r
c = sheet.cell(row = end_r+1, column = 2)
c.value = 'Total'

# colors
colors = counts['colors']

c = sheet.cell(row = 1, column = start_c)
c.value = 'Pre-alignement'
sheet.merge_cells(start_row=1, start_column= start_c, end_row=2, end_column=start_c+len(colors)-1)

for j in range(len(colors)) :
    c = sheet.cell(row = start_r -1, column = j+start_c)
    c.value = colors[j]
    code = colors[j][1:].upper()
    c.fill = openpyxl.styles.PatternFill(start_color = code, fill_type = "solid")

# data
l = counts['per_image_counts']

for row in range(start_r, end_r+1) :
    for col in range(start_c, start_c+len(l[0])) :
        c = sheet.cell(row = row, column = col)
        c.value = l[row-start_r][col-start_c]

# total
for col in range(start_c, start_c+len(colors)) :
    c = sheet.cell(row = end_r +1, column = col)
    c.value = '= SUM('+column_letters[col-1]+str(start_r)+':'+column_letters[col-1]+str(end_r)+')'

excel_wb.save(excel_file+'.xlsx')


# post-alignement
# random correction

images = func.load_slices('slices/'+directories[i]+'4_img_col_correct', nbr_images[i])
counts = func.pixel_count(images)

start_c = start_c + len(colors)

colors = counts['colors']

c = sheet.cell(row = 2, column = start_c)
c.value = 'Random correction'
sheet.merge_cells(start_row=2, start_column= start_c, end_row=2, end_column=start_c+len(colors)-1)

for j in range(len(colors)) :
    c = sheet.cell(row = start_r -1, column = j+start_c)
    c.value = colors[j]
    code = colors[j][1:].upper()
    c.fill = openpyxl.styles.PatternFill(start_color = code, fill_type = "solid")

# data
l = counts['per_image_counts']

for row in range(start_r, end_r+1) :
    for col in range(start_c, start_c+len(l[0])) :
        c = sheet.cell(row = row, column = col)
        c.value = l[row-start_r][col-start_c]

# total
for col in range(start_c, start_c+len(colors)) :
    c = sheet.cell(row = end_r +1, column = col)
    c.value = '= SUM('+column_letters[col-1]+str(start_r)+':'+column_letters[col-1]+str(end_r)+')'

excel_wb.save(excel_file+'.xlsx')

# color classif
images = func.load_slices('slices/'+directories[i]+'4_img_col_classif', nbr_images[i])
counts = func.pixel_count(images)

start_c = start_c + len(colors)

colors = counts['colors']

c = sheet.cell(row = 2, column = start_c)
c.value = 'Color classification'
sheet.merge_cells(start_row=2, start_column= start_c, end_row=2, end_column=start_c+len(colors)-1)

for j in range(len(colors)) :
    c = sheet.cell(row = start_r -1, column = j+start_c)
    c.value = colors[j]
    code = colors[j][1:].upper()
    c.fill = openpyxl.styles.PatternFill(start_color = code, fill_type = "solid")

# data
l = counts['per_image_counts']

for row in range(start_r, end_r+1) :
    for col in range(start_c, start_c+len(l[0])) :
        c = sheet.cell(row = row, column = col)
        c.value = l[row-start_r][col-start_c]

# total
for col in range(start_c, start_c+len(colors)) :
    c = sheet.cell(row = end_r +1, column = col)
    c.value = '= SUM('+column_letters[col-1]+str(start_r)+':'+column_letters[col-1]+str(end_r)+')'

c = sheet.cell(row = 1, column = start_c-len(colors))
#c.value = 'Post-alignement'
#sheet.merge_cells(start_row=1, start_column= start_c-len(colors), end_row=1, end_column=start_c+len(colors)-1)


# images shape
c = sheet.cell(row = 1, column= start_c+len(colors)+1)
c.value = 'Shape of the array images'
c = sheet.cell(row = 2, column= start_c+len(colors)+1)
c.value = str(np.array(images).shape)

excel_wb.save(excel_file+'.xlsx')


